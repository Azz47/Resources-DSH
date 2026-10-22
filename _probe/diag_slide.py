# -*- coding: utf-8 -*-
"""诊断指定页的形状几何与文本，用来定位重排后的重叠 / 断行。

用法: python diag_slide.py <pptx> <页号,页号,...> [--min-w 0.3]
"""
import argparse
import sys

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
EMU = 914400

ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("pages")
ap.add_argument("--min-w", type=float, default=0.0, help="只显示宽度 >= 该值的形状")
a = ap.parse_args()

want = {int(x) for x in a.pages.split(",") if x.strip()}
prs = Presentation(a.src)
SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
print(f"画布 {SW:.2f} x {SH:.2f} in   (2/3 线 = {SW * 2 / 3:.2f})")
print()

for idx, slide in enumerate(prs.slides, 1):
    if idx not in want:
        continue
    print("=" * 100)
    print(f"第 {idx} 页  共 {len(slide.shapes)} 个形状")
    print("=" * 100)
    rows = []
    for sh in slide.shapes:
        try:
            l, t, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
        except Exception:
            rows.append((None, sh, 0, 0, 0, 0))
            continue
        rows.append((sh, sh, l, t, w, h))
    rows.sort(key=lambda r: (round(r[3], 2), round(r[2], 2)))
    for sh, _, l, t, w, h in rows:
        if w < a.min_w:
            continue
        txt = (sh.text_frame.text if getattr(sh, "has_text_frame", False) else "") or ""
        txt = txt.replace("\n", "⏎").replace("\u000b", "⏎")
        kind = sh.shape_type
        anchor = ""
        if getattr(sh, "has_text_frame", False):
            try:
                anchor = str(sh.text_frame.vertical_anchor).split()[0] if sh.text_frame.vertical_anchor is not None else "-"
            except Exception:
                anchor = "-"
            if sh.text_frame.word_wrap is None:
                anchor += "/wrap:None"
        print(f"  x{l:6.2f} y{t:6.2f} w{w:6.2f} h{h:6.2f} | r={l + w:6.2f} b={t + h:6.2f} | {str(kind)[:18]:18s} {anchor:14s} | {txt[:70]}")
    print()
