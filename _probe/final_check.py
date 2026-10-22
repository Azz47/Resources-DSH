# -*- coding: utf-8 -*-
from pptx import Presentation
import zipfile, re
P = r"D:\Resources\DSH\decks\单元思考与行动-参赛版\第二单元 进入新时代-单元思考与行动-参赛版.pptx"
prs = Presentation(P)
notes = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
chars = sum(len(s.notes_slide.notes_text_frame.text.strip()) for s in prs.slides if s.has_notes_slide)
z = zipfile.ZipFile(P)
app = z.read("docProps/app.xml").decode("utf-8", "ignore")
core = z.read("docProps/core.xml").decode("utf-8", "ignore")
print("slides        =", len(prs.slides))
print("notes pages   = %d / %d   讲稿总字数 = %d" % (notes, len(prs.slides), chars))
print("slide size    = %.0f x %.0f pt" % (prs.slide_width / 12700, prs.slide_height / 12700))
print("creator       =", re.search(r"<dc:creator>(.*?)</dc:creator>", core).group(1))
print("lastModifiedBy=", re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core).group(1))
print("Application   =", re.search(r"<Application>(.*?)</Application>", app).group(1))
print("toolprint hit =", [k for k in ("dsh-ppt-studio","pptd","python-pptx") if k in core + app])
media = [n for n in z.namelist() if n.startswith("ppt/media/")]
print("media files   = %d  (%.2f MB)" % (len(media), sum(z.getinfo(n).file_size for n in media) / 1048576))
