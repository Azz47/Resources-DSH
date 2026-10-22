import sys, collections
from pptx import Presentation

def txt(sh):
    if not sh.has_text_frame:
        return ''
    return ' / '.join(p.text for p in sh.text_frame.paragraphs if p.text.strip())

def main(path, label):
    prs = Presentation(path)
    print('=' * 78)
    print(f'{label}  —  字号 → 文本块语义普查')
    print('=' * 78)
    bysize = collections.defaultdict(list)
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip() or not r.font.size:
                        continue
                    sz = r.font.size.pt
                    bysize[sz].append((i, r.text.strip()[:40]))
    for sz in sorted(bysize, reverse=True):
        items = bysize[sz]
        pages = sorted({p for p, _ in items})
        print(f'\n--- {sz}pt : {len(items)} runs / {len(pages)} 页')
        seen = []
        for p, t in items:
            if t not in [s[1] for s in seen]:
                seen.append((p, t))
            if len(seen) >= 6:
                break
        for p, t in seen:
            print(f'    p{p:02d}  {t}')

main(sys.argv[1], sys.argv[2])
