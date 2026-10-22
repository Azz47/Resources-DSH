import json, sys
from pptx import Presentation
from pptx.util import Emu

def geom(path, label):
    prs = Presentation(path)
    print('=' * 76)
    print(f'{label}')
    print(f'  slide size: {prs.slide_width} x {prs.slide_height} EMU '
          f'= {prs.slide_width/914400:.2f}in x {prs.slide_height/914400:.2f}in '
          f'= {prs.slide_width/Emu(1)}x{prs.slide_height/Emu(1)} EMU')
    print(f'  aspect = {prs.slide_width/prs.slide_height:.3f}  (4:3=1.333, 16:9=1.778)')
    print(f'  pages = {len(prs.slides)}')
    sh = []
    for i, sl in enumerate(prs.slides, 1):
        mx = 0
        for s in sl.shapes:
            if s.left is None or s.width is None:
                continue
            mx = max(mx, (s.left + s.width) / Emu(1))
        sh.append((i, mx))
    W = prs.slide_width / Emu(1)
    over = [(i, round(v, 1)) for i, v in sh if v > W + 1]
    print(f'  右边界 max 超过页宽({W:.0f})的页: {len(over)}/{len(sh)} -> {over[:8]}{"..." if len(over) > 8 else ""}')
    body = [(i, round(v, 1)) for i, v in sh if v <= W + 1]
    if body:
        print(f'  未越界的页 右边界: {body[:8]}')

geom(r'D:\Resources\DSH\decks\单元思考与行动-参赛版\第二单元 进入新时代-单元思考与行动-参赛版.pptx', 'OURS-46p')
geom(r'D:\Resources\xwechat_files\wxid_9ka181duelek22_80a9\msg\file\2026-10\单元思考与行动(1).pptx', 'FRIEND-40p')
