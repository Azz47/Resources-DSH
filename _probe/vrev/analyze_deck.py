import sys, json, math
from pptx import Presentation
from pptx.util import Emu

def px(v):
    return None if v is None else round(v / 914400 * 96.0, 1)

def run_size(r, default):
    s = r.font.size
    if s is None:
        return default
    return round(s.pt, 1)

def walk(shapes, out, depth=0):
    for sh in shapes:
        try:
            if sh.shape_type is not None and str(sh.shape_type).startswith('GROUP'):
                walk(sh.shapes, out, depth + 1)
                continue
        except Exception:
            pass
        rec = {
            'name': sh.name,
            'x': px(sh.left), 'y': px(sh.top),
            'w': px(sh.width), 'h': px(sh.height),
        }
        try:
            if sh.has_text_frame:
                tf = sh.text_frame
                txt = tf.text.strip()
                rec['kind'] = 'text'
                rec['text'] = txt[:120]
                rec['tlen'] = len(txt)
                sizes = []
                paras = []
                for p in tf.paragraphs:
                    ptxt = ''.join(r.text for r in p.runs).strip()
                    psz = [run_size(r, None) for r in p.runs if r.text.strip()]
                    psz = [s for s in psz if s]
                    if ptxt:
                        paras.append({'t': ptxt[:80], 'sz': psz[:1]})
                    sizes += psz
                rec['sizes'] = sorted(set(sizes))
                rec['npara'] = len(paras)
                rec['paras'] = paras[:6]
        except Exception:
            pass
        try:
            if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type):
                rec['kind'] = 'image'
        except Exception:
            pass
        if 'kind' not in rec:
            rec['kind'] = 'shape'
        try:
            if sh.fill.type is not None and str(sh.fill.type) == 'SOLID (1)':
                rec['fill'] = str(sh.fill.fore_color.rgb)
        except Exception:
            pass
        out.append(rec)

def main(path):
    prs = Presentation(path)
    W = px(prs.slide_width); H = px(prs.slide_height)
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        shp = []
        walk(slide.shapes, shp)
        pages.append({'page': i, 'shapes': shp,
                      'notes': (slide.notes_slide.notes_text_frame.text.strip()[:200]
                                if slide.has_notes_slide else '')})
    print(json.dumps({'path': path, 'w': W, 'h': H, 'n': len(pages), 'pages': pages},
                     ensure_ascii=False))

if __name__ == '__main__':
    main(sys.argv[1])
