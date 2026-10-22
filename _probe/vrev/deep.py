import sys, collections
from pptx import Presentation
from pptx.util import Emu

def rgb_of(color):
    try:
        if color is not None and color.type is not None:
            return tuple(color.rgb)
    except Exception:
        pass
    return None

def hexs(c): return '#' + ''.join(f'{x:02X}' for x in c)

prs = Presentation(r'D:\Resources\DSH\decks\单元思考与行动-参赛版\第二单元 进入新时代-单元思考与行动-参赛版.pptx')
targets = {24, 37}
for i, sl in enumerate(prs.slides, 1):
    if i not in targets: continue
    print('=' * 70); print(f'PAGE {i}')
    # all solid rects sorted by area desc
    rects = []
    for sh in sl.shapes:
        try:
            if sh.fill.type == 1:
                c = rgb_of(sh.fill.fore_color)
                if c: rects.append(((sh.width or 0)*(sh.height or 0), c, sh))
        except Exception: pass
    rects.sort(key=lambda t: -t[0])
    print('  实心矩形（面积降序）:')
    for a, c, sh in rects[:8]:
        print(f'    {hexs(c)} area={a/914400**2:7.3f} sqin  y={sh.top/914400:.2f}in h={(sh.height or 0)/914400:.2f}in  "{sh.text_frame.text[:24] if sh.has_text_frame else ""}"')
    print('  文字 run（含颜色/字号/所在形状填充）:')
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        filler = None
        try:
            if sh.fill.type == 1: filler = rgb_of(sh.fill.fore_color)
        except Exception: pass
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if not r.text.strip(): continue
                fg = rgb_of(r.font.color)
                sz = r.font.size.pt if r.font.size else None
                print(f'    {hexs(fg) if fg else "(inherit)":>10} {str(sz)+"pt":>8} bg={hexs(filler) if filler else "-":>10} : {r.text[:36]}')
