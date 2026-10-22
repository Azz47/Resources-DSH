import sys, collections
from pptx import Presentation
from pptx.util import Emu

def luma(rgb):
    def f(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = rgb
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)

def contrast(c1, c2):
    a, b = luma(c1), luma(c2)
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)

def solid_fill(sh):
    try:
        f = sh.fill
        if f.type is not None and f.type == 1:  # MSO_FILL.SOLID
            c = f.fore_color
            if c.type is not None and str(c.type).startswith('MSO_THEME_COLOR') is False:
                try:
                    return tuple(c.rgb)
                except Exception:
                    return None
    except Exception:
        pass
    return None

def analyze(path, label):
    prs = Presentation(path)
    print('=' * 78)
    print(label)
    print(f'  slide {prs.slide_width/Emu(1):.0f} x {prs.slide_height/Emu(1):.0f} EMU '
          f'= {prs.slide_width/914400:.2f}in x {prs.slide_height/914400:.2f}in')

    # canvas background
    print('\n  --- 版面/母版 背景填充')
    for sl in list(prs.slides)[:1]:
        lay = sl.slide_layout
        mas = lay.slide_master
        for nm, obj in (('layout', lay), ('master', mas)):
            c = solid_fill(obj)
            print(f'    {nm}: {"#" + "".join(f"{x:02X}" for x in c) if c else "inherit/none"}')

    # per-slide: biggest solid rect = page background; check all text runs contrast
    print('\n  --- 逐页文字对比度（对同页最大实心矩形背景）')
    bad = collections.Counter()
    bad_pages = collections.defaultdict(set)
    total = 0
    for i, sl in enumerate(prs.slides, 1):
        # candidate backgrounds
        rects = []
        for sh in sl.shapes:
            c = solid_fill(sh)
            if c is None:
                continue
            try:
                area = (sh.width or 0) * (sh.height or 0)
            except Exception:
                area = 0
            rects.append((area, c, sh))
        rects.sort(key=lambda t: -t[0])
        bg = rects[0][1] if rects else (255, 255, 255)
        # text runs with explicit colour
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    col = None
                    try:
                        if r.font.color and r.font.color.type is not None:
                            col = tuple(r.font.color.rgb)
                    except Exception:
                        col = None
                    if col is None:
                        continue
                    total += 1
                    cr = contrast(col, bg)
                    sz = r.font.size.pt if r.font.size else 0
                    need = 3.0 if sz >= 18 else 4.5
                    if cr < need:
                        key = ('#' + ''.join(f'{x:02X}' for x in col),
                               '#' + ''.join(f'{x:02X}' for x in bg))
                        bad[key] += 1
                        bad_pages[key].add(i)
    print(f'    显式着色 run 总数 = {total}')
    print(f'    对比度不达标的 (文字色, 背景色) 组合:')
    for (tc, bc), n in bad.most_common(14):
        cr = contrast(tuple(int(tc[j:j+2], 16) for j in (1, 3, 5)),
                      tuple(int(bc[j:j+2], 16) for j in (1, 3, 5)))
        pgs = sorted(bad_pages[(tc, bc)])
        print(f'      {tc} on {bc}  contrast={cr:5.2f}  {n:>4} runs / {len(pgs):>2} 页  {pgs[:14]}{"..." if len(pgs) > 14 else ""}')
    tot_bad = sum(bad.values())
    print(f'    不达标合计 {tot_bad} runs = {tot_bad/max(total,1):.1%}')

    # colour richness
    cols = collections.Counter()
    for sl in prs.slides:
        for sh in sl.shapes:
            c = solid_fill(sh)
            if c:
                cols['#' + ''.join(f'{x:02X}' for x in c)] += 1
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        try:
                            if r.font.color and r.font.color.type is not None:
                                cols['#' + ''.join(f'{x:02X}' for x in tuple(r.font.color.rgb))] += 1
                        except Exception:
                            pass
    print(f'\n  --- 色彩丰富度: 去重颜色数 = {len(cols)}')
    print('    高频:', ', '.join(f'{k}×{v}' for k, v in cols.most_common(12)))

analyze(sys.argv[1], sys.argv[2])
