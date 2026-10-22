# -*- coding: utf-8 -*-
"""读取参赛 PPT 的结构与版式 —— 紧凑输出，不倾倒正文。
重点：每页的**内容右边界**，为"右侧空出 1/3 给讲师站位"做准备。

用法: python inspect_deck.py <pptx> [--full]
"""
import sys
from collections import Counter
from pptx import Presentation
from pptx.util import Emu

EMU = 914400
p = sys.argv[1]
full = "--full" in sys.argv

prs = Presentation(p)
SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
print(f"幻灯片尺寸: {SW:.3f} x {SH:.3f} in  ({'16:9' if abs(SW/SH-16/9)<0.02 else ('4:3' if abs(SW/SH-4/3)<0.02 else '其他')})")
print(f"页数: {len(prs.slides._sldIdLst)}")
print(f"安全区建议: 左侧起 {SW/12:.2f} in；右侧 1/3 从 {SW*2/3:.2f} in 起留空\n")

fonts, colors = Counter(), Counter()
tot_img = 0
rows = []

for idx, slide in enumerate(prs.slides, 1):
    x0, y0, x1, y1 = SW, SH, 0, 0
    n_shapes = n_txt = n_img = n_tbl = 0
    texts = []
    for sh in slide.shapes:
        n_shapes += 1
        try:
            l, t = sh.left / EMU, sh.top / EMU
            w, h = sh.width / EMU, sh.height / EMU
            if w > 0 and h > 0:
                x0, y0 = min(x0, l), min(y0, t)
                x1, y1 = max(x1, l + w), max(y1, t + h)
        except Exception:
            pass
        if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
            n_img += 1
            tot_img += 1
        if getattr(sh, "has_table", False) and sh.has_table:
            n_tbl += 1
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            n_txt += 1
            texts.append(sh.text_frame.text.strip().replace("\n", " / "))
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.name:
                        fonts[r.font.name] += 1
                    try:
                        if r.font.color and r.font.color.rgb:
                            colors[str(r.font.color.rgb)] += 1
                    except Exception:
                        pass
    right_edge = x1 if x1 > 0 else 0
    rows.append((idx, x0, right_edge, n_shapes, n_txt, n_img, n_tbl, texts))

print(f"{'页':>3} {'左起':>6} {'右至':>6} {'形状':>4} {'文本':>4} {'图':>3} {'表':>3}  文本预览")
print("-" * 108)
for idx, x0, x1, ns, nt, ni, ntb, texts in rows:
    preview = " ｜ ".join(texts)[:52] if texts else "(无文字)"
    flag = "  ←右超2/3" if x1 > SW * 2 / 3 else ""
    print(f"{idx:>3} {x0:>6.2f} {x1:>6.2f} {ns:>4} {nt:>4} {ni:>3} {ntb:>3}  {preview}{flag}")

print("\n--- 字体 ---")
for k, v in fonts.most_common(8):
    print(f"   {k:28s} {v}")
print("--- 文字色 ---")
for k, v in colors.most_common(10):
    print(f"   #{k:26s} {v}")
print(f"\n图片总数: {tot_img}")
over = [r for r in rows if r[2] > SW * 2 / 3]
print(f"内容右边界超过 2/3（{SW*2/3:.2f} in）的页: {len(over)} / {len(rows)}")
if over:
    print("   页码: " + ", ".join(str(r[0]) for r in over))

if full:
    print("\n===== 完整文本 =====")
    for idx, *_rest in rows:
        print(f"\n--- P{idx} ---")
        for t in _rest[-1]:
            print("   ", t)
