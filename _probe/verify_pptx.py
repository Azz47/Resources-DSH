# -*- coding: utf-8 -*-
import os, hashlib, sys
from pptx import Presentation
from pptx.util import Emu

cands = [
    r"D:\Resources\xwechat_files\wxid_9ka181duelek22_80a9\msg\file\2026-10\单元思考与行动.pptx",
    r"D:\Resources\DSH\decks\单元思考与行动\source.pptx",
    r"D:\Resources\DSH\decks\单元思考与行动-参赛版\第二单元 进入新时代-单元思考与行动-参赛版.pptx",
    r"D:\Resources\DSH\contest-rework\单元思考与行动_左2-3.pptx",
    r"D:\Resources\DSH\contest-rework\单元思考与行动_右1-3留空.pptx",
    r"D:\Resources\DSH\daofa-demo\八上_第四课第二框_以礼待人\八上_第四课第二框_以礼待人_课件.pptx",
    r"D:\Resources\DSH\daofa-demo\八上_第四课第二框_以礼待人\八上_第四课第二框_以礼待人_课件_参赛版.pptx",
]

def sha(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()[:16]

for p in cands:
    if not os.path.exists(p):
        print(f"MISSING  {p}")
        continue
    st = os.stat(p)
    try:
        pr = Presentation(p)
        n = len(pr.slides)
        w = pr.slide_width; h = pr.slide_height
        # first slide text snippet
        first = ""
        for sh in pr.slides[0].shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                first = sh.text_frame.text.strip().replace("\n", " / ")[:60]
                break
        notes = sum(1 for s in pr.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
        info = f"slides={n} size={w}x{h} pt=({Emu(w).pt:.0f}x{Emu(h).pt:.0f}) noted={notes} first={first}"
    except Exception as e:
        info = f"PARSE_FAIL {e}"
    print(f"{st.st_size:>10} {sha(p)}  {info}\n           {p}")
