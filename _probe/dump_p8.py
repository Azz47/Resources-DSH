# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu

p = r"D:\Resources\DSH\decks\单元思考与行动-参赛版\第二单元 进入新时代-单元思考与行动-参赛版.pptx"
pr = Presentation(p)
s = pr.slides[7]  # page 8
print("=== page 8 shapes (pt) ===")
for sh in s.shapes:
    l = sh.left if sh.left is not None else -1
    t = sh.top if sh.top is not None else -1
    w = sh.width or 0
    h = sh.height or 0
    txt = ""
    if sh.has_text_frame:
        txt = sh.text_frame.text.strip().replace("\n", " / ")[:48]
    wrap = None
    try:
        wrap = sh.text_frame.word_wrap
    except Exception:
        pass
    print(f"  L={Emu(l).pt:6.1f} T={Emu(t).pt:6.1f} W={Emu(w).pt:7.1f} H={Emu(h).pt:6.1f} R={Emu(l+w).pt:7.1f} wrap={wrap} | {txt}")
