# -*- coding: utf-8 -*-
"""诊断为什么某个文本框在缩放后断行：看它的段落/软换行/字号/宽度。

用法: python diag_text.py <pptx> <页号> [,<页号>...]
"""
import sys

from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8")

EMU = 914400.0
prs = Presentation(sys.argv[1])
pages = [int(x) for x in sys.argv[2].split(",")]

for pno in pages:
    sl = prs.slides[pno - 1]
    print("=" * 70)
    print("第 %d 页" % pno)
    for i, sh in enumerate(sl.shapes):
        if not getattr(sh, "has_text_frame", False):
            continue
        tf = sh.text_frame
        txt = tf.text
        if not txt.strip():
            continue
        l, t = sh.left / EMU, sh.top / EMU
        w, h = sh.width / EMU, sh.height / EMU
        print("-" * 70)
        print("  #%d  %s  x%.2f y%.2f w%.2f h%.2f" % (i, sh.shape_type, l, t, w, h))
        print("      全文 repr: %r" % txt)
        print("      段落数 %d  行数(split \\n) %d" % (len(tf.paragraphs), len(txt.split("\n"))))
        for pi, para in enumerate(tf.paragraphs):
            runs = [(r.text, r.font.size.pt if r.font.size else None) for r in para.runs]
            print("      段落%d: %r   runs=%r" % (pi, para.text, runs))
            # 段落级属性，看有没有缩进/边距把可用宽度吃掉
            pPr = para._p.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")
            if pPr is not None:
                print("        pPr attrs: %r" % dict(pPr.attrib))
        bodyPr = tf._txBody.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
        if bodyPr is not None:
            print("      bodyPr attrs: %r" % dict(bodyPr.attrib))
