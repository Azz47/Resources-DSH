# -*- coding: utf-8 -*-
"""v3：把 PPT 重排到「左侧 2/3」——只压横向；保持原有纵向布局；按列避让；字号尽量不动。

v2 的三个错误（已修）：
  1. 把内容强行顶到页顶 → 封面标题撞页眉
  2. 文本框宽度变窄后没跟着长高 → 标题断行压到下个元素
  3. 没做纵向避让 → 元素互相叠

v3 策略：
  · 横向：把内容包围盒映射进 [left, right]，只压 x
  · 纵向：**保留原始 top**，不整体位移
  · 文本重排后按估算高度**长高**；同列内若与上一个元素重叠则下推（列之间互不影响）
  · 若最终底部超界，按档位降字号重算，取能塞下的最大字号

用法: python reflow_v3.py <src.pptx> <out-dir> [--right 8.75]
"""
import argparse
import math
import os

from pptx import Presentation
from pptx.util import Emu

EMU = 914400


def est_h(text, width_in, pt, spacing=1.35):
    if not text or not text.strip():
        return 0.0
    char_w = pt / 72.0
    cpl = max(1, int(width_in / char_w))
    n = 0
    for para in text.replace("\u000b", "\n").split("\n"):
        wide = sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in para)
        n += max(1, math.ceil(wide / cpl))
    return n * (pt / 72.0) * spacing


def max_pt(sh, default=18.0):
    s = []
    if getattr(sh, "has_text_frame", False):
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    s.append(r.font.size.pt)
    return max(s) if s else default


def is_bg(sh, SW, SH):
    try:
        l, w, h = sh.left / EMU, sh.width / EMU, sh.height / EMU
    except Exception:
        return False
    txt = getattr(sh, "has_text_frame", False) and bool(sh.text_frame.text.strip())
    return (not txt) and w >= 0.85 * SW and h >= 0.85 * SH and l <= 0.08 * SW


ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("out_dir")
ap.add_argument("--left", type=float, default=0.30)
ap.add_argument("--right", type=float, default=8.75)
ap.add_argument("--bottom", type=float, default=7.30)
ap.add_argument("--min-font-scale", type=float, default=0.80)
a = ap.parse_args()
os.makedirs(a.out_dir, exist_ok=True)

prs = Presentation(a.src)
SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
LIM = SW * 2 / 3

report = []
for idx, slide in enumerate(prs.slides, 1):
    content = [sh for sh in slide.shapes if not is_bg(sh, SW, SH)]
    geo = []
    for sh in content:
        try:
            geo.append((sh, sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU))
        except Exception:
            pass
    if not geo:
        continue
    x0 = min(g[1] for g in geo)
    x1 = max(g[1] + g[3] for g in geo)
    kx = min(1.0, (a.right - a.left) / max(0.01, x1 - x0))

    # ---- 先算横向压缩后的几何，再按列做纵向避让；字号档位从大到小试 ----
    def layout(fs):
        items = []
        for sh, l, t, w, h in geo:
            nl = (l - x0) * kx + a.left
            nw = max(0.05, w * kx)
            pt = max_pt(sh) * fs
            txt = sh.text_frame.text if getattr(sh, "has_text_frame", False) else ""
            is_pic = sh.__class__.__name__ == "Picture" or sh.shape_type == 13
            nh = (h * kx) if is_pic else max(h, est_h(txt, nw, pt))
            items.append([sh, nl, t, nw, nh])
        # 分列：按 x 区间重叠聚类
        cols = []
        for it in sorted(items, key=lambda z: z[1]):
            for c in cols:
                c_x0 = min(z[1] for z in c)
                c_x1 = max(z[1] + z[3] for z in c)
                if not (it[1] + it[3] <= c_x0 + 0.01 or it[1] >= c_x1 - 0.01):
                    c.append(it)
                    break
            else:
                cols.append([it])
        # 列内按 top 排序后依次避让
        bottom = 0.0
        for c in cols:
            c.sort(key=lambda z: z[2])
            prev_b = None
            for it in c:
                if prev_b is not None and it[2] < prev_b - 0.01:
                    it[2] = prev_b + 0.04
                prev_b = it[2] + it[4]
                bottom = max(bottom, prev_b)
        return items, bottom

    chosen, bottom = None, None
    for fs in [1.0, 0.96, 0.93, 0.90, 0.87, 0.84, a.min_font_scale]:
        if fs < a.min_font_scale - 1e-9:
            continue
        it, b = layout(fs)
        if b <= a.bottom:
            chosen, bottom = it, b
            break
    if chosen is None:
        chosen, bottom = layout(a.min_font_scale)

    for sh, nl, nt, nw, nh in chosen:
        sh.left = Emu(int(nl * EMU))
        sh.top = Emu(int(nt * EMU))
        sh.width = Emu(int(nw * EMU))
        if nh > 0:
            sh.height = Emu(int(nh * EMU))
        if fs != 1.0 and getattr(sh, "has_text_frame", False):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        r.font.size = r.font.size * fs

    report.append((idx, round(kx, 3), round(fs, 3), round(bottom, 2)))

print(f"{'页':>3} {'横向kx':>8} {'字号':>6} {'底部':>6}")
for r in report:
    warn = "  ←超界" if r[3] > a.bottom + 0.02 else ""
    print(f"{r[0]:>3} {r[1]:>8} {r[2]:>6} {r[3]:>6}{warn}")
fss = sorted(r[2] for r in report)
print(f"\n字号中位 {fss[len(fss)//2]:.3f}（v1=0.66 全缩；v2=0.92；v3 目标更接近 1.0）")

out = os.path.join(a.out_dir, os.path.basename(a.src).replace(".pptx", "_左2-3_v3.pptx"))
prs.save(out)
print(f"输出: {out}")
chk = Presentation(out)
bad = [(i, round(max((s2.left + s2.width) / EMU for s2 in s.shapes), 2))
       for i, s in enumerate(chk.slides, 1)
       if max((s2.left + s2.width) / EMU for s2 in s.shapes) > LIM + 0.02]
print(f"右边界超 {LIM:.2f} in 的页: {bad if bad else '无'}")
