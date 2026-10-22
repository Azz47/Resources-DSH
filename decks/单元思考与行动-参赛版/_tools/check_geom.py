# -*- coding: utf-8 -*-
"""体检导出的 pptx：页面尺寸、形状坐标范围、字号分布、备注页数。
用来判断「内容是否真的被放大/裁切」。"""
import glob
import os
import sys
from collections import Counter

from pptx import Presentation

d = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
cands = [p for p in glob.glob(os.path.join(d, "*.pptx")) if "backup" not in os.path.basename(p)]
if len(sys.argv) > 1:
    cands = [sys.argv[1]]

for p in cands:
    prs = Presentation(p)
    W, H = prs.slide_width, prs.slide_height
    print("%s\n  slide = %.0f x %.0f EMU = %.1f x %.1f pt  (%.2f x %.2f in)  slides=%d"
          % (os.path.basename(p), W, H, W / 12700.0, H / 12700.0,
             W / 914400.0, H / 914400.0, len(prs.slides)))
    sizes = Counter()
    notes = 0
    for i, s in enumerate(prs.slides, 1):
        maxr, maxb, minl, mint = -1e9, -1e9, 1e9, 1e9
        for sh in s.shapes:
            try:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
            except Exception:
                continue
            if l is None:
                continue
            minl = min(minl, l / 12700.0)
            mint = min(mint, t / 12700.0)
            maxr = max(maxr, (l + (w or 0)) / 12700.0)
            maxb = max(maxb, (t + (h or 0)) / 12700.0)
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.size:
                            sizes[r.font.size.pt] += 1
        if i <= 3 or i in (14, 15, 16, 46):
            print("   p%-3d x[%.1f, %.1f] y[%.1f, %.1f]" % (i, minl, maxr, mint, maxb))
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            notes += 1
    print("  字号分布(pt):", dict(sorted(sizes.items())))
    print("  带备注页数: %d" % notes)
