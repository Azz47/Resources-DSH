# -*- coding: utf-8 -*-
"""提取 source.pptx 里所有表格的完整内容（导入器把表格丢成了空形状）。"""
import sys
from pptx import Presentation

prs = Presentation(sys.argv[1])
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_table:
            t = sh.table
            print("=== slide %d  table %s  %d rows x %d cols  bounds=(%.0f,%.0f,%.0f,%.0f)" % (
                i, sh.name, len(t.rows), len(t.columns),
                sh.left / 12700.0, sh.top / 12700.0, sh.width / 12700.0, sh.height / 12700.0))
            for r, row in enumerate(t.rows):
                cells = []
                for c in row.cells:
                    txt = (c.text or "").replace("\n", "⏎").strip()
                    cells.append(txt)
                print("  R%02d | %s" % (r, " ¦ ".join(cells)))
            print()
