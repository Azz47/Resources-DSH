import sys, json, io, os, zipfile
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

src = r"D:\Resources\xwechat_files\wxid_9ka181duelek22_80a9\msg\file\2026-10\单元思考与行动(1).pptx"
out = r"D:\Resources\DSH\_wx_analysis"

def pt(v):
    try: return round(Emu(int(v)).pt, 2)
    except Exception: return None

def color_of(fmt):
    try:
        if fmt.type is None: return None
        if str(fmt.type).startswith('SOLID') or fmt.type == 1:
            c = fmt.fore_color
            try: return "#%s" % c.rgb
            except Exception: return "theme:%s" % c.theme_color
        return str(fmt.type)
    except Exception:
        return None

prs = Presentation(src)
print("slide size pt:", pt(prs.slide_width), pt(prs.slide_height), "slides:", len(prs.slides))
print("masters:", len(prs.slide_masters))

def walk(shapes, depth=0):
    res=[]
    for sh in shapes:
        d = {"name": sh.name, "type": str(sh.shape_type), "left": pt(sh.left), "top": pt(sh.top),
             "w": pt(sh.width), "h": pt(sh.height), "depth": depth}
        try: d["rot"]=sh.rotation
        except Exception: pass
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            d["children"]=walk(sh.shapes, depth+1)
        if sh.has_text_frame:
            t = sh.text_frame.text
            d["text"]=t
            runs=[]
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    rr={"t":r.text,"sz":(r.font.size.pt if r.font.size else None),
                        "b":r.font.bold, "name":r.font.name}
                    try: rr["col"]="#%s"%r.font.color.rgb
                    except Exception:
                        try: rr["col"]="theme:%s"%r.font.color.theme_color
                        except Exception: rr["col"]=None
                    runs.append(rr)
            d["runs"]=runs
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                img = sh.image
                d["img"] = {"ext":img.ext, "px":[img.size[0],img.size[1]], "bytes":len(img.blob), "sha":img.sha1[:10]}
            except Exception as e:
                d["img_err"]=str(e)
        try:
            if sh.fill.type is not None:
                d["fill"]=color_of(sh.fill)
        except Exception: pass
        try:
            if sh.line.fill.type is not None:
                d["line"]=color_of(sh.line.fill)
                d["line_w"]=sh.line.width.pt if sh.line.width else None
        except Exception: pass
        try:
            if sh.is_placeholder: d["ph"]=str(sh.placeholder_format.type)
        except Exception: pass
        res.append(d)
    return res

slides=[]
for i, s in enumerate(prs.slides, 1):
    bg=None
    try:
        bg = color_of(s.background.fill)
    except Exception: pass
    notes = ""
    try:
        if s.has_notes_slide: notes = s.notes_slide.notes_text_frame.text
    except Exception: pass
    slides.append({"idx":i, "bg":bg, "notes_len":len(notes), "notes":notes[:400], "shapes":walk(s.shapes)})

json.dump({"slides":slides}, open(os.path.join(out,"wx_dump.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=1)
print("dumped")

# per-page quick geometry summary
for sl in slides:
    def flat(ss):
        for s in ss:
            yield s
            if "children" in s:
                yield from flat(s["children"])
    allsh=list(flat(sl["shapes"]))
    rights=[ (s["left"] or 0)+(s["w"] or 0) for s in allsh if s["left"] is not None and s["w"] is not None]
    bots=[ (s["top"] or 0)+(s["h"] or 0) for s in allsh if s["top"] is not None and s["h"] is not None]
    npic=sum(1 for s in allsh if s["type"].endswith("PICTURE (13)"))
    txtlen=sum(len(s.get("text","")) for s in allsh)
    print(f"p{sl['idx']:02d} bg={sl['bg']} shapes={len(allsh):3d} pic={npic} maxR={max(rights) if rights else 0:7.1f} maxB={max(bots) if bots else 0:6.1f} chars={txtlen} notes={sl['notes_len']}")
